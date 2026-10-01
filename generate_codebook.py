# -*- coding: utf-8 -*-
"""
BDEF 编码法 · 码书生成器
=========================================================================
让每个字符获得 1 ~ 30030 之间的唯一序号，生成可逆的"码书"。

输出（默认写到脚本所在目录）：
    codebook.txt   制表符分隔的可读码书：编码 <TAB> 字符 <TAB> Unicode <TAB> 类别 <TAB> UTF-8
    codebook.xlsx  带表头、冻结窗格的 Excel 版码书
    codebook.js    网页版使用的码书（window.BDEF_CODEBOOK = "……"）

依赖：
    openpyxl（仅 codebook.xlsx 需要；若未安装会自动跳过，不影响 txt / js）
        pip install openpyxl

用法：
    python generate_codebook.py
    python generate_codebook.py --outdir ./dist --no-xlsx

原理备注：
    模数 11、13、14、15 两两互质，乘积 11*13*14*15 = 30030。
    编码 n（1<=n<=30030）-> 四个余数 (n%11, n%13, n%14, n%15) -> 十六进制拼接 = 4 位编码，
    由中国剩余定理可唯一还原。本脚本只负责"字符 <-> 序号"这张表。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import unicodedata

# --------------------------------------------------------------------------
# 可调参数
# --------------------------------------------------------------------------
TARGET_TOTAL = 30030      # 码书总条目数
HANZI_START = 0x4E00      # 常用汉字起始码点（一）
HANZI_END = 0x9FFF        # 汉字区结束码点；覆盖完整 CJK 统一汉字区，简繁通吃
HANZI_COUNT = HANZI_END - HANZI_START + 1   # = 20992
# 说明：原文《公歧子：BDEF编码法》取前 15000 个汉字（到 U+8897）。
#      若需与原文完全对齐，运行 `python generate_codebook.py --hanzi-count 15000`。
#      两种口径下编码 1~15000（一、丁、丂 … ）完全相同。

# --------------------------------------------------------------------------
# 收集器
# --------------------------------------------------------------------------


class Codebook:
    """按顺序收集字符，自动去重、过滤不可用字符。"""

    hanzi_count: int = HANZI_COUNT      # 开头收录的汉字数量

    def __init__(self, target: int = TARGET_TOTAL) -> None:
        self.target = target
        self.chars: list[tuple[str, str]] = []   # (字符, 类别名)
        self._seen: set[str] = set()

    # -- 基础设施 ----------------------------------------------------------
    @property
    def size(self) -> int:
        return len(self.chars)

    @property
    def full(self) -> bool:
        return len(self.chars) >= self.target

    @staticmethod
    def _acceptable(ch: str) -> bool:
        """过滤控制字符、代理区、未分配字符、格式字符等。"""
        if not isinstance(ch, str) or len(ch) != 1:
            return False
        code = ord(ch)
        if code < 0x20:                                   # C0 控制字符
            return False
        if 0x7F <= code <= 0x9F:                          # DEL / C1 控制字符
            return False
        if 0xD800 <= code <= 0xDFFF:                      # 代理区
            return False
        if code in (0xFFFE, 0xFFFF, 0x2028, 0x2029):      # 非字符 / 行分隔符
            return False
        if ch in "\r\n\t":
            return False
        if unicodedata.category(ch) in ("Cc", "Cs", "Cn", "Cf"):
            return False
        return True

    def add_char(self, ch: str, category: str) -> bool:
        if self.full:
            return False
        if ch in self._seen:
            return False
        if not self._acceptable(ch):
            return False
        self._seen.add(ch)
        self.chars.append((ch, category))
        return True

    # -- 批量收录 ----------------------------------------------------------
    def add_range(self, start: int, end: int, category: str) -> int:
        """收录 [start, end] 区间内所有可用码点，返回新增数量。"""
        n = 0
        for cp in range(start, end + 1):
            if self.full:
                break
            if self.add_char(chr(cp), category):
                n += 1
        return n

    def add_string(self, text: str, category: str) -> int:
        n = 0
        for ch in text:
            if self.full:
                break
            if self.add_char(ch, category):
                n += 1
        return n

    # -- 各类字符 ----------------------------------------------------------
    def add_common_hanzi(self) -> int:
        """常用中文字：从 U+4E00 起连续取 self.hanzi_count 个。"""
        n = self.hanzi_count
        return self.add_range(HANZI_START, HANZI_START + n - 1,
                              "常用中文字")

    def add_english_latin(self) -> int:
        """英文字母大小写 + ASCII 可打印 + 拉丁扩展各族。"""
        n = 0
        n += self.add_string("ABCDEFGHIJKLMNOPQRSTUVWXYZ", "英文字母")
        n += self.add_string("abcdefghijklmnopqrstuvwxyz", "英文字母")
        n += self.add_string("".join(chr(c) for c in range(0x20, 0x7F)),
                             "ASCII 可打印")
        n += self.add_range(0x00A0, 0x00FF, "拉丁补充")        # Latin-1 Supplement
        n += self.add_range(0x0100, 0x017F, "拉丁扩展-A")      # Latin Extended-A
        n += self.add_range(0x0180, 0x024F, "拉丁扩展-B")      # Latin Extended-B
        n += self.add_range(0x0300, 0x036F, "组合附加符号")    # Combining Diacritics
        return n

    def add_japanese(self) -> int:
        """日文：平假名 / 片假名 / 半角片假名。"""
        n = 0
        n += self.add_range(0x3040, 0x309F, "日文平假名")
        n += self.add_range(0x30A0, 0x30FF, "日文片假名")
        n += self.add_range(0xFF65, 0xFF9F, "日文半角片假名")
        return n

    def add_punctuation(self) -> int:
        """标点：中文标点 / 全角符号 / 通用标点。"""
        n = 0
        n += self.add_range(0x3000, 0x303F, "中文标点")
        n += self.add_range(0xFF00, 0xFF64, "全角符号")
        n += self.add_range(0x2000, 0x206F, "通用标点")
        return n

    def add_symbols(self) -> int:
        """符号：上下标 / 货币 / 箭头 / 数学运算 / 几何图形 / 装饰符号。"""
        n = 0
        n += self.add_range(0x2070, 0x209F, "上标下标")
        n += self.add_range(0x20A0, 0x20CF, "货币符号")
        n += self.add_range(0x2190, 0x21FF, "箭头符号")
        n += self.add_range(0x2200, 0x22FF, "数学运算符")
        n += self.add_range(0x25A0, 0x25FF, "几何图形")
        n += self.add_range(0x2700, 0x27BF, "装饰符号")
        n += self.add_range(0x2460, 0x24FF, "带圈字符")
        return n

    def add_greek_cyrillic(self) -> int:
        """希腊字母 + 西里尔字母。"""
        n = 0
        n += self.add_range(0x0370, 0x03FF, "希腊字母")
        n += self.add_range(0x0400, 0x04FF, "西里尔字母")
        return n

    def add_korean(self) -> int:
        """韩文：谚文字母 / 兼容字母 / 谚文音节。"""
        n = 0
        n += self.add_range(0x3130, 0x318F, "韩文兼容字母")
        n += self.add_range(0x1100, 0x11FF, "韩文谚文字母")
        n += self.add_range(0xAC00, 0xAC00 + 4999, "韩文音节")
        return n

    def add_thai_vietnamese(self) -> int:
        """泰文 + 越南文拉丁扩展附加。"""
        n = 0
        n += self.add_range(0x0E00, 0x0E7F, "泰文")
        n += self.add_range(0x1EA0, 0x1EFF, "越南文拉丁扩展附加")
        n += self.add_range(0x1E00, 0x1E9F, "拉丁扩展附加")
        return n

    def fill_to_target(self) -> int:
        """补齐到 TARGET_TOTAL：谚文音节 -> 汉字 -> CJK 扩展 A -> 扩展 B -> 私有区。"""
        n = 0
        plan = (
            (0xAD00, 0xD7A3, "韩文音节"),
            (0x9FA6, 0x9FFF, "汉字"),
            (0x3400, 0x4DBF, "CJK 扩展 A"),
            (0x20000, 0x2A6DF, "CJK 扩展 B"),
            (0xE000, 0xF8FF, "私有区占位"),
            (0xF900, 0xFAFF, "CJK 兼容汉字"),
            (0xAC00, 0xACFF, "韩文音节"),
        )
        for start, end, cat in plan:
            if self.full:
                break
            n += self.add_range(start, end, cat)
        # 极端兜底：用私有区逐个补满
        cp = 0xE000
        while not self.full:
            n += self.add_char(chr(cp), "私有区占位")
            cp += 1
            if cp > 0xF8FF:
                break
        return n


def build_codebook(target: int = TARGET_TOTAL, hanzi_count: int = HANZI_COUNT
                   ) -> list[tuple[str, str]]:
    """按既定顺序构建码书，返回 [(字符, 类别), ...]，长度 == target。"""
    cb = Codebook(target)
    cb.hanzi_count = hanzi_count
    cb.add_common_hanzi()
    cb.add_english_latin()
    cb.add_japanese()
    cb.add_punctuation()
    cb.add_symbols()
    cb.add_greek_cyrillic()
    cb.add_korean()
    cb.add_thai_vietnamese()
    cb.fill_to_target()
    return cb.chars


# --------------------------------------------------------------------------
# 输出
# --------------------------------------------------------------------------


def save_txt(rows: list[tuple[str, str]], path: str) -> None:
    """制表符分隔的纯文本码书。"""
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("# BDEF 编码法 · 码书（codebook）\n")
        f.write("# 格式：编码<TAB>字符<TAB>Unicode<TAB>类别<TAB>UTF-8\n")
        f.write("# 共 %d 条，编码范围 1 ~ %d\n" % (len(rows), len(rows)))
        for i, (ch, cat) in enumerate(rows, start=1):
            utf8 = " ".join("%02X" % b for b in ch.encode("utf-8"))
            f.write("%d\t%s\tU+%04X\t%s\t%s\n" % (i, ch, ord(ch), cat, utf8))


def save_js(rows: list[tuple[str, str]], path: str) -> None:
    """网页版使用的码书：window.BDEF_CODEBOOK = "<30030 个字符>";

    额外导出类别分组，便于网页展示统计。
    """
    text = "".join(ch for ch, _ in rows)
    literal = json.dumps(text, ensure_ascii=False)
    # 统计各类别数量
    groups: "list[list]" = []
    for ch, cat in rows:
        if groups and groups[-1][0] == cat:
            groups[-1][1] += 1
        else:
            groups.append([cat, 1])
    banner = (
        "/* BDEF 编码法 · 码书（自动生成，请勿手工修改）\n"
        " * 由 generate_codebook.py 生成，共 %d 个字符。\n"
        " * window.BDEF_CODEBOOK 中第 i 个字符的编码即为 i+1。\n"
        " */\n" % len(rows)
    )
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(banner)
        f.write("window.BDEF_CODEBOOK_GROUPS = %s;\n"
                % json.dumps(groups, ensure_ascii=False))
        f.write("window.BDEF_CODEBOOK = %s;\n" % literal)


def save_xlsx(rows: list[tuple[str, str]], path: str) -> bool:
    """Excel 版码书，返回是否成功。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  [跳过] 未安装 openpyxl，不生成 xlsx（pip install openpyxl）")
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "码书"
    header = ["编码", "字符", "Unicode", "类别", "UTF-8"]
    ws.append(header)

    head_font = Font(bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", fgColor="4F46E5")
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, (ch, cat) in enumerate(rows, start=1):
        utf8 = " ".join("%02X" % b for b in ch.encode("utf-8"))
        ws.append([i, ch, "U+%04X" % ord(ch), cat, utf8])
        # 以 "=" 开头的字符强制为文本，避免 Excel 当公式处理
        ws.cell(row=i + 1, column=2).data_type = "s"

    for col, width in zip(range(1, 6), (10, 10, 12, 24, 26)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    wb.save(path)
    return True


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------


def main(argv: "list[str] | None" = None) -> int:
    parser = argparse.ArgumentParser(
        description="生成 BDEF 编码法码书（codebook）")
    parser.add_argument("--outdir", default=os.path.dirname(os.path.abspath(__file__)),
                        help="输出目录，默认脚本所在目录")
    parser.add_argument("--total", type=int, default=TARGET_TOTAL,
                        help="码书总条目数，默认 30030")
    parser.add_argument("--hanzi-count", type=int, default=HANZI_COUNT,
                        dest="hanzi",
                        help="开头收录的汉字数量，默认 %d（完整 CJK 统一汉字区）；"
                             "填 15000 可复现原文口径" % HANZI_COUNT)
    parser.add_argument("--no-xlsx", action="store_true", help="不生成 xlsx")
    args = parser.parse_args(argv)

    os.makedirs(args.outdir, exist_ok=True)
    rows = build_codebook(args.total, args.hanzi)

    txt_path = os.path.join(args.outdir, "codebook.txt")
    js_path = os.path.join(args.outdir, "codebook.js")
    xlsx_path = os.path.join(args.outdir, "codebook.xlsx")

    print("实际生成条目数：%d" % len(rows))
    if len(rows) != args.total:
        print("  [警告] 期望 %d 条，实际 %d 条" % (args.total, len(rows)))

    save_txt(rows, txt_path)
    print("已保存：%s" % txt_path)
    save_js(rows, js_path)
    print("已保存：%s" % js_path)
    if not args.no_xlsx:
        if save_xlsx(rows, xlsx_path):
            print("已保存：%s" % xlsx_path)

    print("\n前 10 条预览：")
    for i, (ch, cat) in enumerate(rows[:10], start=1):
        print("  %2d  %s  U+%04X  %s" % (i, ch, ord(ch), cat))
    return 0


if __name__ == "__main__":
    sys.exit(main())
